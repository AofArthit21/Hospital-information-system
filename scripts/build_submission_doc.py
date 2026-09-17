from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
DIAGRAM_PATH = ROOT / "docs" / "ER_DIAGRAM.png"
DOC_PATH = OUTPUT_DIR / "Agnos_Backend_Assignment_Documentation.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_index % 2:
                shade(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], str(value))
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F6F8FA")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])


def load_font(path, size):
    try:
        return ImageFont.truetype(str(path), size)
    except OSError:
        return ImageFont.load_default()


def build_er_diagram():
    image = Image.new("RGB", (1800, 1040), "white")
    draw = ImageDraw.Draw(image)
    fonts = Path("C:/Windows/Fonts")
    title_font = load_font(fonts / "arialbd.ttf", 44)
    header_font = load_font(fonts / "arialbd.ttf", 28)
    body_font = load_font(fonts / "arial.ttf", 24)
    label_font = load_font(fonts / "arialbd.ttf", 23)

    draw.text((60, 35), "Hospital Middleware — Entity Relationship Diagram", fill="#1F2937", font=title_font)

    def entity_box(x, y, width, title, fields):
        row_height = 42
        header_height = 64
        height = header_height + row_height * len(fields) + 18
        draw.rounded_rectangle((x, y, x + width, y + height), radius=14, fill="#FFFFFF", outline="#1F4E78", width=4)
        draw.rounded_rectangle((x, y, x + width, y + header_height), radius=14, fill="#1F4E78", outline="#1F4E78")
        draw.rectangle((x, y + header_height - 15, x + width, y + header_height), fill="#1F4E78")
        draw.text((x + 22, y + 14), title, fill="white", font=header_font)
        current_y = y + header_height + 10
        for index, (key, name, field_type) in enumerate(fields):
            if index % 2:
                draw.rectangle((x + 3, current_y - 2, x + width - 3, current_y + row_height - 3), fill="#F4F8FB")
            draw.text((x + 18, current_y + 4), key, fill="#B45309", font=label_font)
            draw.text((x + 90, current_y + 4), name, fill="#111827", font=body_font)
            text_width = draw.textlength(field_type, font=body_font)
            draw.text((x + width - text_width - 18, current_y + 4), field_type, fill="#6B7280", font=body_font)
            current_y += row_height
        return (x, y, x + width, y + height)

    hospital = entity_box(80, 150, 520, "HOSPITALS", [
        ("PK", "id", "bigint"), ("UK", "code", "varchar(50)"), ("", "name", "varchar(255)"),
        ("", "his_base_url", "text"), ("", "created_at", "timestamptz"), ("", "updated_at", "timestamptz"),
    ])
    staff = entity_box(80, 600, 520, "STAFF", [
        ("PK", "id", "bigint"), ("FK", "hospital_id", "bigint"), ("", "username", "varchar(100)"),
        ("", "password_hash", "text"), ("", "created_at", "timestamptz"), ("", "updated_at", "timestamptz"),
    ])
    patient = entity_box(990, 150, 720, "PATIENTS", [
        ("PK", "id", "bigint"), ("FK", "hospital_id", "bigint"), ("UK", "patient_hn", "varchar(100)"),
        ("UK", "national_id", "varchar(30)"), ("UK", "passport_id", "varchar(30)"),
        ("", "first/middle/last_name_th", "varchar(100)"), ("", "first/middle/last_name_en", "varchar(100)"),
        ("", "date_of_birth", "date"), ("", "phone_number", "varchar(30)"), ("", "email", "varchar(255)"),
        ("", "gender", "M | F"), ("", "created_at / updated_at", "timestamptz"),
    ])

    def relation(start, end, label):
        draw.line((start[0], start[1], end[0], end[1]), fill="#2563EB", width=5)
        draw.polygon([(end[0], end[1]), (end[0] - 20, end[1] - 12), (end[0] - 20, end[1] + 12)], fill="#2563EB")
        middle_x = (start[0] + end[0]) // 2
        middle_y = (start[1] + end[1]) // 2
        draw.rounded_rectangle((middle_x - 100, middle_y - 24, middle_x + 100, middle_y + 24), 10, fill="white", outline="#93C5FD")
        text_width = draw.textlength(label, font=label_font)
        draw.text((middle_x - text_width / 2, middle_y - 15), label, fill="#1D4ED8", font=label_font)

    relation((hospital[2], 285), (patient[0], 285), "1   owns   N")
    relation((340, hospital[3]), (340, staff[1]), "1 employs N")
    draw.text((80, 980), "Every staff and patient row belongs to one hospital. All patient queries are scoped by hospital_id from the JWT.", fill="#374151", font=body_font)
    DIAGRAM_PATH.parent.mkdir(parents=True, exist_ok=True)
    image.save(DIAGRAM_PATH, quality=95)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.color.rgb = RGBColor.from_string(BLUE)
    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(BLUE)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Agnos Candidate Assignment")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Backend Developer — Development Planning & Technical Documentation").bold = True
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("Go • Gin • PostgreSQL • Docker Compose • Nginx\nVersion 1.0 • September 2026")
    document.add_paragraph()
    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.add_run("A hospital middleware API for authenticated, hospital-scoped patient discovery and HIS integration.")
    document.add_page_break()

    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(
        "This implementation provides APIs for hospital staff registration and authentication, retrieves patient records from a hospital information system (HIS), normalizes and caches those records in PostgreSQL, and supports multi-field patient search. The authenticated staff member's hospital is derived from the signed JWT and applied to every patient query, preventing cross-hospital data access."
    )
    add_table(document, ["Requirement", "Implementation"], [
        ("Backend", "Go 1.23 with Gin HTTP framework"),
        ("Database", "PostgreSQL 16 with constraints, indexes, and seed migration"),
        ("Authentication", "bcrypt password hashes and expiring HS256 JWT access tokens"),
        ("HIS integration", "Timeout-bound HTTP adapter for GET /patient/search/{id}"),
        ("Deployment", "Docker Compose with private API/PostgreSQL network and Nginx ingress"),
        ("Quality", "Unit, regression, OpenAPI syntax, and Docker runtime acceptance tests"),
    ], [1.5, 5.5])

    document.add_heading("2. Architecture and Project Structure", level=1)
    document.add_heading("2.1 Request Flow", level=2)
    add_bullets(document, [
        "Staff registration resolves the hospital, hashes the password with bcrypt, and stores a hospital-scoped account.",
        "Login verifies hospital, username, and password, then returns a signed JWT containing staff_id and hospital_id.",
        "JWT middleware validates the signature and expiry and places the trusted hospital_id in request context.",
        "An ID search calls the configured HIS, validates the payload, and upserts the normalized patient record.",
        "The repository searches PostgreSQL with hospital_id as the mandatory first predicate and adds only parameterized filters.",
    ])
    document.add_heading("2.2 Package Layout", level=2)
    add_code(document, """cmd/api/                         Application entry point and graceful shutdown
internal/auth/                   JWT creation and validation
internal/config/                 Environment configuration
internal/domain/                 Entities, filters, and domain errors
internal/his/                    Hospital A HTTP adapter
internal/httpapi/                Gin routes, middleware, handlers, response mapping
internal/repository/postgres/    Parameterized PostgreSQL repositories
internal/service/                Authentication and patient use cases
migrations/                      Schema, constraints, indexes, and seed data
deploy/                          Nginx reverse-proxy configuration
test/e2e/                        Deterministic mock HIS fixture
docs/                            OpenAPI, ER diagram, and implementation plan""")
    document.add_heading("2.3 Component Responsibilities", level=2)
    add_table(document, ["Layer", "Responsibility", "Reason"], [
        ("HTTP", "Validation, status codes, DTO mapping", "Keeps transport concerns out of business rules"),
        ("Service", "Authentication and patient workflows", "Use cases can be unit-tested with interfaces"),
        ("Repository", "SQL persistence and hospital-scoped queries", "Centralizes data ownership enforcement"),
        ("HIS adapter", "External HTTP request and response normalization", "Isolates upstream failure behavior"),
        ("Domain", "Entities, filters, stable errors", "Avoids framework coupling"),
    ], [1.1, 3.0, 2.9])

    document.add_heading("3. API Specification", level=1)
    document.add_paragraph("Base URL: http://localhost:8080. The complete machine-readable contract is available in docs/openapi.yaml (OpenAPI 3.0.3).")
    add_table(document, ["Method", "Path", "Authentication", "Purpose"], [
        ("GET", "/healthz", "None", "Container liveness check"),
        ("POST", "/staff/create", "None", "Create hospital staff login credentials"),
        ("POST", "/staff/login", "None", "Authenticate staff and issue JWT"),
        ("GET", "/patient/search", "Bearer JWT", "Search patients in the authenticated hospital"),
    ], [0.7, 1.7, 1.2, 3.4])

    document.add_heading("3.1 Create Staff", level=2)
    add_code(document, """POST /staff/create
Content-Type: application/json

{
  "username": "alice",
  "password": "password123",
  "hospital": "hospital-a"
}""")
    add_table(document, ["Status", "Meaning"], [("201", "Staff created"), ("400", "Invalid request or unknown hospital"), ("409", "Username already exists in the hospital")], [1.0, 6.0])

    document.add_heading("3.2 Staff Login", level=2)
    add_code(document, """POST /staff/login
Content-Type: application/json

{
  "username": "alice",
  "password": "password123",
  "hospital": "hospital-a"
}

200 OK
{
  "access_token": "<JWT>",
  "token_type": "Bearer",
  "expires_in": 28800
}""")

    document.add_heading("3.3 Patient Search", level=2)
    document.add_paragraph("All query parameters are optional. With no parameters, the API returns up to 100 most recently updated patients in the caller's hospital.")
    add_table(document, ["Parameter", "Type", "Matching behavior"], [
        ("national_id", "string", "Exact; refreshes from HIS first"),
        ("passport_id", "string", "Exact; refreshes from HIS first"),
        ("first_name", "string", "Case-insensitive partial match across Thai/English"),
        ("middle_name", "string", "Case-insensitive partial match across Thai/English"),
        ("last_name", "string", "Case-insensitive partial match across Thai/English"),
        ("date_of_birth", "YYYY-MM-DD", "Exact date"),
        ("phone_number", "string", "Exact"),
        ("email", "email", "Case-insensitive exact match"),
    ], [1.5, 1.2, 4.3])
    add_code(document, """GET /patient/search?national_id=1234567890123
Authorization: Bearer <access_token>

200 OK
{
  "data": [{
    "first_name_th": "อลิส",
    "middle_name_th": "",
    "last_name_th": "ทดสอบ",
    "first_name_en": "Alice",
    "middle_name_en": "",
    "last_name_en": "Test",
    "date_of_birth": "1990-01-02",
    "patient_hn": "HN-A-001",
    "national_id": "1234567890123",
    "phone_number": "0812345678",
    "email": "alice@example.com",
    "gender": "F"
  }],
  "count": 1
}""")
    document.add_heading("3.4 Error Contract", level=2)
    add_code(document, """{
  "error": {
    "code": "VALIDATION_ERROR",
    "message": "one or more search parameters are invalid"
  }
}""")
    add_table(document, ["HTTP", "Code", "When"], [
        ("400", "VALIDATION_ERROR / UNKNOWN_HOSPITAL", "Request validation or hospital lookup fails"),
        ("401", "UNAUTHORIZED / INVALID_CREDENTIALS", "Missing/invalid token or login fails"),
        ("409", "STAFF_EXISTS", "Duplicate hospital-scoped username"),
        ("500", "INTERNAL_ERROR", "Unexpected internal failure without leaked details"),
        ("502", "HIS_UNAVAILABLE", "HIS timeout, 5xx, or invalid upstream payload"),
    ], [0.7, 2.3, 4.0])

    document.add_page_break()
    document.add_heading("4. Database Schema and ER Diagram", level=1)
    document.add_picture(str(DIAGRAM_PATH), width=Inches(7.0))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Figure 1. Hospital ownership model and patient/staff relationships.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("4.1 Constraints and Indexes", level=2)
    add_bullets(document, [
        "hospitals.code is globally unique and lowercase.",
        "staff (hospital_id, username) is unique.",
        "patients (hospital_id, patient_hn) is the unique HIS cache/upsert key.",
        "Non-null national_id and passport_id are individually unique within each hospital.",
        "Every patient requires at least one national or passport identifier; gender is M or F when present.",
        "Hospital/name and hospital/date-of-birth indexes support common search patterns.",
    ])

    document.add_heading("5. Docker Compose Deployment", level=1)
    add_code(document, """Client
  │
  ▼
Nginx :8080  ──private network──>  Go/Gin API :8080  ──>  PostgreSQL :5432
                                          │
                                          └────────────> Hospital HIS HTTPS API""")
    add_table(document, ["Service", "Image/Build", "Exposure", "Health"], [
        ("nginx", "nginx:1.27-alpine", "Host port 8080", "Proxies only to API"),
        ("api", "Multi-stage Go build; distroless runtime", "Private network", "GET /healthz"),
        ("postgres", "postgres:16-alpine", "Private network", "pg_isready"),
    ], [1.0, 2.5, 1.5, 2.0])
    add_code(document, """docker compose up -d --build
docker compose ps
docker compose logs -f api nginx
docker compose down""")

    document.add_heading("6. Security and Privacy", level=1)
    add_bullets(document, [
        "Hospital scope comes from the validated JWT, never from patient-search input.",
        "Passwords are bcrypt-hashed and are never returned or logged.",
        "SQL uses positional parameters; no user input is concatenated into queries.",
        "HIS calls have a configurable timeout and a 1 MiB response-body limit.",
        "Gin and Nginx access logs deliberately omit query strings containing patient identifiers.",
        "The API image runs as the distroless nonroot user, and only Nginx publishes a host port.",
        "Nginx limits request rate and body size; API responses do not leak internal errors.",
    ])

    document.add_heading("7. Testing and Verification", level=1)
    add_table(document, ["Test layer", "Coverage"], [
        ("HTTP handlers", "Positive/negative staff creation, login, authorization, validation, HIS failure mapping"),
        ("Services", "bcrypt, token claims, HIS cache workflow, hospital scope propagation"),
        ("HIS adapter", "Valid response normalization, not found, invalid payload"),
        ("Regression", "Dynamic SQL repeated placeholders and query-string log privacy"),
        ("Runtime acceptance", "14 Docker API checks through Nginx, including cross-hospital isolation"),
    ], [1.5, 5.5])
    add_code(document, """go test -count=1 ./...
go vet ./...
docker compose config --quiet
docker compose up -d --build""")
    document.add_paragraph("Verified runtime result: API healthy, nonroot runtime, zero restarts, PostgreSQL healthy, and 14/14 acceptance checks passed through Nginx.")

    document.add_heading("8. Evaluation Criteria Mapping", level=1)
    add_table(document, ["Criterion", "Evidence"], [
        ("Requirement satisfaction", "All requested APIs, models, authentication, hospital isolation, Docker services, and tests are implemented."),
        ("Code quality", "Layered packages, interface boundaries, graceful shutdown, parameterized SQL, stable domain errors, and regression tests."),
        ("Unit test coverage", "Positive and negative scenarios for each API plus service, adapter, and regression tests."),
        ("Documentation clarity", "This document, OpenAPI 3.0.3 specification, README, implementation plan, and ER diagram."),
    ], [1.7, 5.3])

    document.add_heading("9. Repository Entry Points", level=1)
    add_bullets(document, [
        "README.md — setup instructions, curl examples, behavior, and production notes",
        "docs/openapi.yaml — machine-readable API contract",
        "docs/ER_DIAGRAM.md and docs/ER_DIAGRAM.png — schema relationships",
        "docs/PLAN.md — architecture decisions and test strategy",
        "docker-compose.yml — Nginx, Go API, and PostgreSQL topology",
        "migrations/001_init.sql — complete schema and Hospital A seed",
    ])

    for current_section in document.sections:
        footer = current_section.footer.paragraphs[0]
        footer.add_run("Agnos Backend Candidate Assignment  |  Confidential  |  ")
        add_page_number(footer)

    document.save(DOC_PATH)


if __name__ == "__main__":
    build_er_diagram()
    build_document()
    print(DOC_PATH)
    print(DIAGRAM_PATH)
